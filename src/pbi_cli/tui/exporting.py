"""Take what the Explorer shows out of the terminal: onto the clipboard, or into a file.

The Explorer shows tables (the people with access to a report, the stored versions of a list,
the items of a workspace) and the JSON the lake holds. What a person wants to do with them is
paste them into a mail or a spreadsheet, or keep them as a file to share. This module holds what
a tab can give (`Exportable`) and the formats it can be written in. It imports nothing of
Textual, so a script can use it too.

```python
from pbi_cli.tui.exporting import Exportable, write

people = Exportable("Users of Sales", ("Name", "E-mail"), (("Ann", "ann@example.com"),))
write(people, "csv", "users-of-sales.csv")
```

A cell that a spreadsheet would take for a formula (it starts with ``=``, ``+``, ``-`` or
``@``) is written as text, so that opening an export never runs anything: in CSV and tab-separated
text by a leading ``'``, in a workbook as a text cell.
"""

import csv
import io
import json
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Dict, List, Sequence, Tuple, Union

CSV = "csv"
XLSX = "xlsx"
JSON = "json"
MARKDOWN = "md"
TSV = "tsv"

#: The formats of a table, the first being the usual one, and of the JSON the lake holds.
TABLE_FORMATS = (CSV, XLSX, JSON, MARKDOWN, TSV)
DATA_FORMATS = (JSON,)

#: What each format is, in words.
TITLES = {
    CSV: "CSV (opens in Excel)",
    XLSX: "Excel workbook",
    JSON: "JSON",
    MARKDOWN: "Markdown table",
    TSV: "Tab-separated text",
}

EXTENSIONS = {CSV: ".csv", XLSX: ".xlsx", JSON: ".json", MARKDOWN: ".md", TSV: ".tsv"}

#: A cell that starts with one of these is a formula to a spreadsheet.
_FORMULA = ("=", "+", "-", "@", "\t", "\r")

#: What a table of the Explorer shows for a cell that has nothing in it.
_EMPTY = ("-", "–")

_UNSAFE_IN_NAMES = re.compile(r'[<>:"/\\|?*\x00-\x1f]')
_RESERVED_NAMES = {"con", "prn", "aux", "nul"} | {
    f"{device}{number}" for device in ("com", "lpt") for number in range(1, 10)
}

#: A name of a worksheet: 31 characters, none of ``[]:*?/\``.
_SHEET_LIMIT = 31
_UNSAFE_IN_SHEETS = re.compile(r"[\[\]:*?/\\]")


@dataclass(frozen=True)
class Exportable:
    """What a tab of the Explorer can give.

    :param title: what it is, for the name of the file and the notice (``Users of Sales``)
    :param columns: the columns of a table (none: it is not a table)
    :param rows: the rows of a table, each a tuple of texts as long as the columns
    :param data: for what is not a table, the JSON that the lake holds
    """

    title: str
    columns: Tuple[str, ...] = ()
    rows: Tuple[Tuple[str, ...], ...] = ()
    data: Any = None

    @property
    def is_table(self) -> bool:
        return bool(self.columns)

    @property
    def formats(self) -> Tuple[str, ...]:
        """The formats it can be written in, the usual one first."""
        return TABLE_FORMATS if self.is_table else DATA_FORMATS

    @property
    def headers(self) -> Tuple[str, ...]:
        """The columns with a name each: a blank one is ``column 2``, and a name that is
        there twice gets a number."""
        names: List[str] = []
        for position, column in enumerate(self.columns, start=1):
            name = column.strip() or f"column {position}"
            candidate, number = name, 1
            while candidate in names:
                number += 1
                candidate = f"{name} {number}"
            names.append(candidate)
        return tuple(names)

    @property
    def cells(self) -> List[Tuple[str, ...]]:
        """The rows with what shows as ``-`` in a table of the Explorer made empty."""
        return [
            tuple("" if cell.strip() in _EMPTY else cell for cell in row)
            for row in self.rows
        ]

    @property
    def count(self) -> int:
        """How many rows it has (1 for the JSON of a single thing)."""
        return len(self.rows) if self.is_table else 1

    def describe(self) -> str:
        """What it is, in words: ``Users of Sales (12 rows)``."""
        if not self.is_table:
            return self.title
        return f"{self.title} ({self.count} row{'' if self.count == 1 else 's'})"


def _text_cell(value: str) -> str:
    """A cell for a format that a spreadsheet reads as text it may take for a formula."""
    return "'" + value if value.startswith(_FORMULA) else value


def _objects(item: Exportable) -> List[Dict[str, str]]:
    names = item.headers
    return [dict(zip(names, row)) for row in item.cells]


def to_text(item: Exportable, fmt: str) -> str:
    """The text of an export, for every format that is text (not ``xlsx``).

    :raises ValueError: for a format that this cannot be written in
    """
    if fmt not in item.formats:
        raise ValueError(f"{item.title} cannot be written as {fmt}")
    if not item.is_table:
        return json.dumps(item.data, indent=2, ensure_ascii=False, default=str) + "\n"
    if fmt == JSON:
        return json.dumps(_objects(item), indent=2, ensure_ascii=False) + "\n"
    if fmt == CSV:
        out = io.StringIO()
        writer = csv.writer(out, lineterminator="\r\n")
        writer.writerow([_text_cell(name) for name in item.headers])
        for row in item.cells:
            writer.writerow([_text_cell(cell) for cell in row])
        return out.getvalue()
    if fmt == TSV:
        lines = [item.headers] + item.cells
        return (
            "\n".join(
                "\t".join(_text_cell(re.sub(r"[\t\r\n]+", " ", c)) for c in row)
                for row in lines
            )
            + "\n"
        )
    if fmt == MARKDOWN:

        def cell(value: str) -> str:
            return value.replace("|", "\\|").replace("\r", "").replace("\n", "<br>")

        head = "| " + " | ".join(cell(name) for name in item.headers) + " |"
        rule = "| " + " | ".join("---" for _ in item.headers) + " |"
        body = ["| " + " | ".join(cell(c) for c in row) + " |" for row in item.cells]
        return "\n".join([head, rule, *body]) + "\n"
    raise ValueError(f"{fmt} is not a format of text")  # pragma: no cover


def clipboard_text(item: Exportable) -> str:
    """What to put on the clipboard: a table as tab-separated text (it pastes into the cells
    of a spreadsheet), anything else as JSON."""
    return to_text(item, TSV if item.is_table else JSON)


def _legal(value: str) -> str:
    """A text without the control characters that a workbook cannot hold (a vertical tab, a
    form feed, ...): openpyxl refuses them."""
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

    return ILLEGAL_CHARACTERS_RE.sub("", value)


def sheet_name(title: str) -> str:
    """A name for the worksheet of a workbook."""
    name = _UNSAFE_IN_SHEETS.sub(" ", _legal(title)).strip("' ")
    return (name[:_SHEET_LIMIT] or "Export").strip()


def _write_workbook(item: Exportable, path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    book = Workbook()
    sheet = book.active
    sheet.title = sheet_name(item.title)
    widths = [len(name) for name in item.headers]
    for position, name in enumerate(item.headers, start=1):
        cell = sheet.cell(row=1, column=position, value=_legal(name))
        cell.font = Font(bold=True)
    for number, row in enumerate(item.cells, start=2):
        for position, value in enumerate(row, start=1):
            cell = sheet.cell(row=number, column=position, value=_legal(value))
            cell.data_type = "s"  # text, even when it starts with "=": never a formula
            widths[position - 1] = max(widths[position - 1], len(value))
    for position, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(position)].width = min(width + 2, 60)
    sheet.freeze_panes = "A2"
    if item.rows:
        sheet.auto_filter.ref = sheet.dimensions
    book.save(path)


def write(item: Exportable, fmt: str, path: Union[str, Path]) -> Path:
    """Write an export to a file, making the folder if it is not there.

    CSV and tab-separated text are written in UTF-8 with a byte order mark, which is what
    Excel needs to read accents and other letters right.

    :return: the path written
    :raises ValueError: for a format that the item cannot be written in
    :raises OSError: when the file cannot be written
    """
    if fmt not in item.formats:
        raise ValueError(f"{item.title} cannot be written as {fmt}")
    target = Path(path).expanduser()
    target.parent.mkdir(parents=True, exist_ok=True)
    if fmt == XLSX:
        _write_workbook(item, target)
    else:
        encoding = "utf-8-sig" if fmt in (CSV, TSV) else "utf-8"
        target.write_text(to_text(item, fmt), encoding=encoding, newline="")
    return target


def file_stem(title: str, limit: int = 80) -> str:
    """A name for a file from a title: no character that a system refuses in a name, no name
    that Windows keeps for a device, and not too long."""
    stem = _UNSAFE_IN_NAMES.sub("-", re.sub(r"\s+", " ", title))
    stem = stem.strip(" .-")[:limit].rstrip(" .-")
    if not stem:
        stem = "export"
    if stem.split(".")[0].strip().lower() in _RESERVED_NAMES:
        stem = f"_{stem}"
    return stem


def file_name(item: Exportable, fmt: str) -> str:
    """The name that an export is suggested to have: ``Users of Sales.csv``."""
    return file_stem(item.title) + EXTENSIONS[fmt]


def free_path(folder: Union[str, Path], name: str) -> Path:
    """A path in a folder that no file has yet: the name, or the name with ``(2)``, ``(3)``
    ... before its extension."""
    base = Path(folder).expanduser()
    candidate = base / name
    stem, suffix = Path(name).stem, Path(name).suffix
    number = 1
    while candidate.exists():
        number += 1
        candidate = base / f"{stem} ({number}){suffix}"
    return candidate


def rows_of(table: Sequence[Sequence[Any]]) -> Tuple[Tuple[str, ...], ...]:
    """Rows of anything that shows as text, as tuples of texts."""
    return tuple(
        tuple("" if cell is None else str(cell) for cell in row) for row in table
    )
