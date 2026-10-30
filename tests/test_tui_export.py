"""Taking what the Explorer shows out of it: ``y`` copies it, ``x`` saves it or copies it."""

import asyncio
import csv
import io
import json

import pytest
from openpyxl import load_workbook
from sync_helpers import World
from textual.widgets import Input, RadioButton, RadioSet
from tui_helpers import backend_of, plain, run_ui

from pbi_cli.core.scan import ScanFlags
from pbi_cli.core.store import LakeStore
from pbi_cli.tui.commands import commands_for
from pbi_cli.tui.exporting import Exportable
from pbi_cli.tui.modals import ExportModal


@pytest.fixture
def world(tmp_path):
    world = World(tmp_path)
    world.run()
    world.run(
        "scan", scan_flags=ScanFlags.everything(), workspace_ids=("ws-0001", "ws-0002")
    )
    world.run("group-users", force=True)
    world.run("report-users", force=True)
    world.run("app-users", force=True)
    return world


@pytest.fixture
def folder(tmp_path, monkeypatch):
    """The folder `pbi tui` was started in."""
    place = tmp_path / "work"
    place.mkdir()
    monkeypatch.chdir(place)
    return place


def notes(ui):
    return [n.message for n in ui.app._notifications]


async def on_the_users_of_a_report(ui, tab="2"):
    await ui.select("workspace", "ws-0001")
    await ui.pick_row("report:rep-0001")
    await ui.press(tab)


def dialog(ui):
    screen = ui.app.screen
    assert isinstance(screen, ExportModal)
    return screen


async def set_path(ui, text):
    dialog(ui).query_one("#path", Input).value = str(text)
    await ui.settle()


def path_of(ui):
    return dialog(ui).query_one("#path", Input).value


def error_of(ui):
    return ui.static("#error", dialog(ui))


def formats_on_offer(ui):
    return [
        b.name for b in dialog(ui).query_one("#format", RadioSet).query(RadioButton)
    ]


def note_of(ui):
    return ui.static("#format-note", dialog(ui))


def format_picked(ui):
    return dialog(ui).query_one("#format", RadioSet).pressed_button.name


async def pick_format(ui, name):
    for button in dialog(ui).query_one("#format", RadioSet).query(RadioButton):
        if button.name == name:
            button.value = True
    await ui.settle()


# ---------------------------------------------------------------------------
# y: copy
# ---------------------------------------------------------------------------


def test_y_copies_the_tab_for_a_spreadsheet(world, clipboard_of_the_test):
    async def scenario(ui):
        await on_the_users_of_a_report(ui)
        await ui.press("y")
        return ui.app.clipboard, notes(ui)

    app_clipboard, said = run_ui(backend_of(world), scenario)

    expected = (
        "Name\tE-mail or id\tAccess\tType\nAnn\tann-rep-0001@example.com\tOwner\tUser\n"
    )
    assert app_clipboard == expected  # for pasting into an input of the app
    assert clipboard_of_the_test == [expected]  # what the tool of the system got
    assert said == ["Copied Users of report Report 1 (1 row) to the clipboard."]


def test_y_copies_the_table_when_the_tab_has_nothing_to_give(
    world, clipboard_of_the_test
):
    async def scenario(ui):
        await ui.select(
            "workspace", "ws-0003"
        )  # nothing is fetched of the people of items
        await ui.press("2")
        await ui.press("y")
        return notes(ui)

    # the people of the workspace are in the lake, so the Users tab gives them
    said = run_ui(backend_of(world), scenario)
    assert said == ["Copied Users of workspace Workspace 3 (2 rows) to the clipboard."]

    async def with_a_tab_that_says_nothing(ui):
        await ui.select("workspace", "ws-0001")
        await ui.pick_row("dataset:ds-0001")
        await ui.press("2")  # the people of a dataset were never fetched
        await ui.press("y")
        return notes(ui)

    said = run_ui(backend_of(world), with_a_tab_that_says_nothing)
    assert said == ["Copied Items of Workspace 1 (5 rows) to the clipboard."]
    assert (
        clipboard_of_the_test[-1].splitlines()[0] == "Name\tType\tOwner\tUpdated\tFrom"
    )


def test_y_says_when_there_is_nothing_to_copy(tmp_path, clipboard_of_the_test):
    empty = World(tmp_path)

    async def scenario(ui):
        ui.explorer.query_one("#table-filter", Input).value = "zzzz"  # no row matches
        await ui.settle()
        await ui.press("2")  # and the people of a lake that holds nothing: none
        await ui.press("y")
        return notes(ui)

    said = run_ui(backend_of(empty), scenario)

    assert said[-1].startswith("There is nothing to export here")
    assert clipboard_of_the_test == []


def test_an_empty_lake_still_has_an_info_tab_to_copy(tmp_path, clipboard_of_the_test):
    empty = World(tmp_path)

    async def scenario(ui):
        await ui.press("y")
        return notes(ui)

    said = run_ui(backend_of(empty), scenario)

    assert said == ["Copied Info of The lake is empty (4 rows) to the clipboard."]


def test_a_terminal_is_asked_when_the_system_has_no_tool(
    world, clipboard_of_the_test, monkeypatch
):
    from pbi_cli.tui import clipboard

    monkeypatch.setattr(clipboard, "system_copy", lambda text: "")

    async def scenario(ui):
        await on_the_users_of_a_report(ui)
        await ui.press("y")
        return ui.app.clipboard, notes(ui)

    app_clipboard, said = run_ui(backend_of(world), scenario)

    assert app_clipboard.startswith("Name\tE-mail")  # the terminal was asked
    assert said[-1].startswith("Sent Users of report Report 1 (1 row) to the clipboard")
    assert "save it to a file instead: press x" in said[-1]


def test_a_tool_of_the_system_that_breaks_does_not_end_the_app(
    world, clipboard_of_the_test, monkeypatch
):
    from pbi_cli.tui import clipboard

    def broken(text):
        raise RuntimeError("no clipboard here")

    monkeypatch.setattr(clipboard, "system_copy", broken)

    async def scenario(ui):
        await on_the_users_of_a_report(ui)
        await ui.press("y")
        return notes(ui)

    assert run_ui(backend_of(world), scenario)[-1].startswith("Sent Users of report")


# ---------------------------------------------------------------------------
# x: the dialog
# ---------------------------------------------------------------------------


def test_x_offers_the_tab_and_the_table_and_suggests_a_file_in_the_folder(
    world, folder
):
    async def scenario(ui):
        await on_the_users_of_a_report(ui)
        await ui.press("x")
        what = [
            plain(b.label)
            for b in dialog(ui).query_one("#what", RadioSet).query(RadioButton)
        ]
        return what, path_of(ui), formats_on_offer(ui), format_picked(ui)

    what, path, formats, picked = run_ui(backend_of(world), scenario)

    assert what == [
        "The Users tab: Users of report Report 1 (1 row)",
        "The table: Items of Workspace 1 (5 rows)",
    ]
    assert path == str(folder / "Users of report Report 1.csv")
    assert formats == ["csv", "xlsx", "json", "md", "tsv"] and picked == "csv"


def test_saving_writes_the_file_and_says_where(world, folder):
    async def scenario(ui):
        await on_the_users_of_a_report(ui)
        await ui.press("x")
        await ui.click("#save")
        return type(ui.app.screen).__name__, notes(ui), ui.app.export_folder

    screen, said, remembered = run_ui(backend_of(world), scenario)

    target = folder / "Users of report Report 1.csv"
    assert screen == "ExplorerScreen"  # the dialog closed
    assert said == [f"Saved Users of report Report 1 (1 row) to {target}"]
    assert remembered == str(folder)
    rows = list(csv.reader(io.StringIO(target.read_text(encoding="utf-8-sig"))))
    assert rows == [
        ["Name", "E-mail or id", "Access", "Type"],
        ["Ann", "ann-rep-0001@example.com", "Owner", "User"],
    ]


def test_enter_in_the_file_name_saves_as_well(world, folder):
    async def scenario(ui):
        await on_the_users_of_a_report(ui)
        await ui.press("x")
        await ui.press("enter")
        return type(ui.app.screen).__name__

    assert run_ui(backend_of(world), scenario) == "ExplorerScreen"
    assert (folder / "Users of report Report 1.csv").exists()


def test_the_next_dialog_suggests_the_folder_of_the_last_file(world, folder, tmp_path):
    elsewhere = tmp_path / "elsewhere" / "deeper"

    async def scenario(ui):
        await on_the_users_of_a_report(ui)
        await ui.press("x")
        await set_path(ui, elsewhere / "mine.csv")
        await ui.click("#save")
        await ui.press("x")
        return path_of(ui)

    assert run_ui(backend_of(world), scenario) == str(
        elsewhere / "Users of report Report 1.csv"
    )
    assert (elsewhere / "mine.csv").exists()  # the folders were made


def test_the_table_can_be_saved_as_a_workbook(world, folder):
    async def scenario(ui):
        await on_the_users_of_a_report(ui)
        await ui.press("x")
        dialog(ui).query_one("#what", RadioSet).query(RadioButton)[1].value = True
        await ui.settle()
        await pick_format(ui, "xlsx")
        suggested = path_of(ui)
        await ui.click("#save")
        return suggested

    suggested = run_ui(backend_of(world), scenario)

    target = folder / "Items of Workspace 1.xlsx"
    assert suggested == str(target)
    sheet = load_workbook(target).active
    assert [c.value for c in sheet[1]] == ["Name", "Type", "Owner", "Updated", "From"]
    assert sheet.max_row == 6  # the header and five items


def test_the_extension_follows_the_format_until_the_name_is_typed(world, folder):
    async def scenario(ui):
        await on_the_users_of_a_report(ui)
        await ui.press("x")
        await pick_format(ui, "json")
        first = path_of(ui)
        await set_path(ui, folder / "mine.csv")
        await pick_format(ui, "md")
        return first, path_of(ui)

    first, kept = run_ui(backend_of(world), scenario)

    assert first == str(folder / "Users of report Report 1.json")
    assert kept == str(folder / "mine.md")  # the name stays, the extension follows


def test_a_name_without_a_known_extension_is_left_as_it_is(world, folder):
    async def scenario(ui):
        await on_the_users_of_a_report(ui)
        await ui.press("x")
        await set_path(ui, folder / "users.txt")
        await pick_format(ui, "json")
        return path_of(ui)

    assert run_ui(backend_of(world), scenario) == str(folder / "users.txt")


def test_a_suggested_name_is_never_one_that_exists(world, folder):
    (folder / "Users of report Report 1.csv").write_text("mine")

    async def scenario(ui):
        await on_the_users_of_a_report(ui)
        await ui.press("x")
        return path_of(ui)

    assert run_ui(backend_of(world), scenario) == str(
        folder / "Users of report Report 1 (2).csv"
    )


def test_a_file_that_exists_is_replaced_only_when_saving_is_asked_for_twice(
    world, folder
):
    target = folder / "mine.csv"
    target.write_text("keep me")

    async def scenario(ui):
        await on_the_users_of_a_report(ui)
        await ui.press("x")
        await set_path(ui, target)
        await ui.click("#save")
        first = error_of(ui), type(ui.app.screen).__name__, target.read_text()
        await asyncio.sleep(
            0.4
        )  # a button ignores a click while it still shows it was pressed
        await ui.click("#save")
        return first, type(ui.app.screen).__name__

    first, screen = run_ui(backend_of(world), scenario)

    assert first[0].startswith(
        "That file exists already: press Save again to replace it"
    )
    assert first[1:] == ("ExportModal", "keep me")
    assert screen == "ExplorerScreen"
    assert "Name,E-mail or id" in target.read_text(encoding="utf-8-sig")


def test_typing_another_name_forgets_that_the_file_exists(world, folder):
    target = folder / "mine.csv"
    target.write_text("keep me")

    async def scenario(ui):
        await on_the_users_of_a_report(ui)
        await ui.press("x")
        await set_path(ui, target)
        await ui.click("#save")
        await set_path(ui, folder / "another.csv")
        return error_of(ui)

    assert run_ui(backend_of(world), scenario) == ""
    assert target.read_text() == "keep me"


@pytest.mark.parametrize(
    "typed, said",
    [
        ("", "Give the file a name."),
        ("   ", "Give the file a name."),
        ("<folder>", "That is a folder: add the name of the file to it."),
    ],
)
def test_a_file_that_cannot_be_named_is_said_and_the_dialog_stays(
    world, folder, typed, said
):
    async def scenario(ui):
        await on_the_users_of_a_report(ui)
        await ui.press("x")
        await set_path(ui, typed.replace("<folder>", str(folder)))
        await ui.click("#save")
        return error_of(ui), type(ui.app.screen).__name__

    assert run_ui(backend_of(world), scenario) == (said, "ExportModal")


def test_a_file_that_cannot_be_written_is_said_and_the_dialog_stays(world, folder):
    blocker = folder / "a-file"
    blocker.write_text("x")  # a folder cannot be made inside a file

    async def scenario(ui):
        await on_the_users_of_a_report(ui)
        await ui.press("x")
        await set_path(ui, blocker / "inside.csv")
        await ui.click("#save")
        return error_of(ui), type(ui.app.screen).__name__

    said, screen = run_ui(backend_of(world), scenario)

    assert said.startswith("Could not write the file: ") and screen == "ExportModal"


def test_quotes_around_a_path_are_ignored(world, folder):
    async def scenario(ui):
        await on_the_users_of_a_report(ui)
        await ui.press("x")
        await set_path(ui, f'"{folder / "quoted.csv"}"')  # as Windows copies a path
        await ui.click("#save")

    run_ui(backend_of(world), scenario)

    assert (folder / "quoted.csv").exists()


def test_cancel_writes_nothing(world, folder):
    async def scenario(ui):
        await on_the_users_of_a_report(ui)
        await ui.press("x")
        await ui.click("#cancel")
        return type(ui.app.screen).__name__, notes(ui)

    screen, said = run_ui(backend_of(world), scenario)

    assert screen == "ExplorerScreen" and said == []
    assert list(folder.iterdir()) == []


def test_escape_closes_the_dialog(world, folder):
    async def scenario(ui):
        await on_the_users_of_a_report(ui)
        await ui.press("x")
        await ui.press("escape")
        return type(ui.app.screen).__name__

    assert run_ui(backend_of(world), scenario) == "ExplorerScreen"


def test_the_copy_button_puts_the_text_of_the_format_on_the_clipboard(
    world, folder, clipboard_of_the_test
):
    async def scenario(ui):
        await on_the_users_of_a_report(ui)
        await ui.press("x")
        await pick_format(ui, "json")
        await ui.click("#copy")
        return notes(ui)

    said = run_ui(backend_of(world), scenario)

    assert json.loads(clipboard_of_the_test[-1]) == [
        {
            "Name": "Ann",
            "E-mail or id": "ann-rep-0001@example.com",
            "Access": "Owner",
            "Type": "User",
        }
    ]
    assert said == ["Copied Users of report Report 1 (1 row) to the clipboard."]
    assert list(folder.iterdir()) == []  # nothing was written


def test_a_workbook_cannot_be_copied_so_the_text_for_cells_is(
    world, folder, clipboard_of_the_test
):
    async def scenario(ui):
        await on_the_users_of_a_report(ui)
        await ui.press("x")
        await pick_format(ui, "xlsx")
        await ui.click("#copy")

    run_ui(backend_of(world), scenario)

    assert clipboard_of_the_test[-1].startswith("Name\tE-mail or id\tAccess\tType\n")


def test_the_note_under_the_formats_says_what_the_picked_one_is(world, folder):
    async def scenario(ui):
        await on_the_users_of_a_report(ui)
        await ui.press("x")
        first = note_of(ui)
        await pick_format(ui, "xlsx")
        return first, note_of(ui)

    assert run_ui(backend_of(world), scenario) == (
        "Opens in Excel (letters such as accents are kept).",
        "An Excel workbook with a header row you can filter.",
    )


def test_a_dialog_that_is_answered_twice_acts_once(
    world, folder, clipboard_of_the_test
):
    async def scenario(ui):
        await on_the_users_of_a_report(ui)
        await ui.press("x")
        screen = dialog(ui)
        screen._copy()
        screen._copy()  # Enter and the button, before the dialog has gone
        await ui.settle()
        return type(ui.app.screen).__name__

    assert run_ui(backend_of(world), scenario) == "ExplorerScreen"
    assert len(clipboard_of_the_test) == 1


def test_the_format_becomes_one_that_the_other_export_can_have(world, tmp_path):
    table = Exportable("A table", ("a",), (("1",),))
    data = Exportable("Some data", data={"x": 1})

    async def scenario(ui):
        ui.app.push_screen(
            ExportModal([("The table", table), ("The JSON", data)], str(tmp_path))
        )
        await ui.settle()
        first = format_picked(ui), path_of(ui)
        dialog(ui).query_one("#what", RadioSet).query(RadioButton)[1].value = True
        await ui.settle()
        return first, (
            format_picked(ui),
            formats_on_offer(ui),
            path_of(ui),
            note_of(ui),
        )

    first, then = run_ui(backend_of(world), scenario)

    assert first == ("csv", str(tmp_path / "A table.csv"))
    assert then == (
        "json",
        ["json"],
        str(tmp_path / "Some data.json"),
        "A list of objects, one for each row.",
    )


def test_the_json_tab_can_only_be_saved_as_json(world, folder):
    async def scenario(ui):
        await on_the_users_of_a_report(ui, tab="4")
        await ui.press("x")
        dialog(ui).query_one("#what", RadioSet)  # the tab and the table are offered
        return path_of(ui), formats_on_offer(ui), format_picked(ui)

    path, formats, picked = run_ui(backend_of(world), scenario)

    assert path == str(folder / "JSON of report Report 1.json")
    assert formats == ["json"] and picked == "json"


def test_choosing_the_table_after_the_json_tab_gives_the_formats_of_a_table_back(
    world, folder
):
    async def scenario(ui):
        await on_the_users_of_a_report(ui, tab="4")
        await ui.press("x")
        dialog(ui).query_one("#what", RadioSet).query(RadioButton)[1].value = True
        await ui.settle()
        return formats_on_offer(ui), format_picked(ui), path_of(ui)

    formats, picked, path = run_ui(backend_of(world), scenario)

    assert formats == ["csv", "xlsx", "json", "md", "tsv"] and picked == "json"
    assert path == str(
        folder / "Items of Workspace 1.json"
    )  # the format that was picked


def test_the_json_that_is_saved_is_what_the_lake_holds(world, folder):
    async def scenario(ui):
        await on_the_users_of_a_report(ui, tab="4")
        await ui.press("x")
        await ui.click("#save")

    run_ui(backend_of(world), scenario)

    saved = json.loads((folder / "JSON of report Report 1.json").read_text("utf-8"))
    assert saved["id"] == "rep-0001" and saved["name"] == "Report 1"


def test_with_only_one_thing_to_export_there_is_no_choice_of_what(
    tmp_path, tmp_path_factory, monkeypatch
):
    monkeypatch.chdir(tmp_path_factory.mktemp("work"))
    empty = World(tmp_path)

    async def scenario(ui):
        ui.explorer.query_one("#table-filter", Input).value = "zzzz"  # no row matches
        await ui.settle()
        await ui.press("x")  # the Info tab is shown: it has fields to give
        screen = dialog(ui)
        return ui.static("#what-one", screen), len(screen.query("#what")), path_of(ui)

    what, choices, path = run_ui(backend_of(empty), scenario)

    assert what == "The Info tab: Info of The lake is empty (4 rows)"
    assert choices == 0 and path.endswith("Info of The lake is empty.csv")


# ---------------------------------------------------------------------------
# what each tab gives
# ---------------------------------------------------------------------------


def test_the_users_of_the_workspaces_are_saved_with_the_workspace_of_each(
    world, folder
):
    async def scenario(ui):
        await ui.select("workspaces")
        await ui.press("2")
        await ui.press("x")
        await ui.click("#save")

    run_ui(backend_of(world), scenario)

    target = folder / "Users of 12 workspaces.csv"
    rows = list(csv.reader(io.StringIO(target.read_text(encoding="utf-8-sig"))))
    assert rows[0] == [
        "Name",
        "E-mail or id",
        "Access",
        "Type",
        "Workspace",
        "Workspace id",
    ]
    assert len(rows) == 1 + 24 and ["Workspace 1", "ws-0001"] == rows[1][4:]


def test_what_the_filter_leaves_in_the_table_is_what_the_table_exports(world, folder):
    async def scenario(ui):
        await ui.select("workspaces")
        ui.explorer.query_one("#table-filter", Input).value = "workspace 12"
        await ui.settle()
        await ui.press("2")
        await ui.press("x")
        await ui.click("#save")

    run_ui(backend_of(world), scenario)

    rows = list(
        csv.reader(
            io.StringIO(
                (folder / "Users of 1 workspace.csv").read_text(encoding="utf-8-sig")
            )
        )
    )
    assert len(rows) == 3 and {r[4] for r in rows[1:]} == {"Workspace 12"}


def test_the_scan_tab_saves_where_each_table_reads_its_data_from(world, folder):
    async def scenario(ui):
        await ui.select("workspace", "ws-0001")
        await ui.press("7")
        await ui.press("x")
        await ui.click("#save")

    run_ui(backend_of(world), scenario)

    rows = list(
        csv.DictReader(
            io.StringIO(
                (folder / "Sources of workspace Workspace 1.csv").read_text(
                    encoding="utf-8-sig"
                )
            )
        )
    )
    sales = next(r for r in rows if r["Table"] == "Sales")
    assert sales["Source"] == "SQL Server" and sales["Server"] == "sql.example"
    assert sales["Item"] == "dbo.Sales" and "Sql.Database" in sales["Expression"]


def test_the_users_of_an_app_picked_in_the_list_of_apps_are_saved(world, folder):
    async def scenario(ui):
        await ui.select("apps")
        await ui.pick_row("app:app-0002")
        await ui.press("2")
        await ui.press("x")
        await ui.click("#save")

    run_ui(backend_of(world), scenario)

    target = folder / "Users of app App 2.csv"
    rows = list(csv.reader(io.StringIO(target.read_text(encoding="utf-8-sig"))))
    assert rows == [
        ["Name", "E-mail or id", "Access", "Type"],
        ["Ann", "ann-app-0002@example.com", "Owner", "User"],
        ["Bob", "bob-app-0002@example.com", "Read", "User"],
    ]


def test_the_freshness_dots_of_the_table_are_saved_as_words(world, folder):
    async def scenario(ui):
        await ui.select("workspaces")
        ui.explorer.query_one("#table-filter", Input).value = "workspace 1"
        await ui.settle()
        await ui.press("1")
        await ui.press("x")
        dialog(ui).query_one("#what", RadioSet).query(RadioButton)[1].value = True
        await ui.settle()
        await ui.click("#save")

    run_ui(backend_of(world), scenario)

    rows = list(
        csv.DictReader(
            io.StringIO((folder / "Workspaces.csv").read_text(encoding="utf-8-sig"))
        )
    )
    assert rows[0]["Freshness"] == "fresh" and rows[0]["Name"] == "Workspace 1"
    assert "●" not in (folder / "Workspaces.csv").read_text(encoding="utf-8-sig")


# ---------------------------------------------------------------------------
# the palette
# ---------------------------------------------------------------------------


def test_the_palette_offers_to_export_and_to_copy(world, folder, clipboard_of_the_test):
    async def scenario(ui):
        await on_the_users_of_a_report(ui)
        found = {c.title: c for c in commands_for(ui.app, ui.app.screen)}
        found["Copy what is shown"].run()
        await ui.settle()
        found["Export what is shown…"].run()
        await ui.settle()
        return (
            found["Export what is shown…"].help,
            found["Copy what is shown"].help,
            type(ui.app.screen).__name__,
        )

    export_help, copy_help, screen = run_ui(backend_of(world), scenario)

    assert export_help.endswith("or copy it  ·  key x")
    assert copy_help.endswith("to paste into a spreadsheet  ·  key y")
    assert screen == "ExportModal" and len(clipboard_of_the_test) == 1


def test_a_lake_that_is_only_looked_at_can_be_exported_too(world, folder):
    async def scenario(ui):
        await on_the_users_of_a_report(ui)
        await ui.press("x")
        await ui.click("#save")

    view_only = backend_of(
        world,
        store=LakeStore(world.store.root, readonly=True, reason="The lake only reads."),
    )
    run_ui(view_only, scenario)

    assert (folder / "Users of report Report 1.csv").exists()
