"""What leaves the Explorer: the formats of an export, the names of its files, its text."""

import csv
import io
import json

import pytest
from openpyxl import load_workbook

from pbi_cli.tui import exporting
from pbi_cli.tui.exporting import Exportable

PEOPLE = Exportable(
    "Users of Sales",
    ("Name", "E-mail or id", "Access", "Type"),
    (
        ("Ann Ås", "ann@example.com", "Admin", "User"),
        ("Bob", "-", "Viewer", "User"),
    ),
)
THING = Exportable("JSON of Sales", data={"id": "w1", "name": "Sales ü", "n": [1, 2]})


# ---------------------------------------------------------------------------
# what can be written how
# ---------------------------------------------------------------------------


def test_a_table_can_be_written_in_five_formats_and_other_data_only_as_json():
    assert PEOPLE.is_table and PEOPLE.formats == ("csv", "xlsx", "json", "md", "tsv")
    assert not THING.is_table and THING.formats == ("json",)


def test_a_format_that_an_export_cannot_have_is_refused(tmp_path):
    with pytest.raises(ValueError, match="cannot be written as csv"):
        exporting.to_text(THING, "csv")
    with pytest.raises(ValueError, match="cannot be written as csv"):
        exporting.write(THING, "csv", tmp_path / "x.csv")
    with pytest.raises(ValueError, match="cannot be written as pdf"):
        exporting.to_text(PEOPLE, "pdf")


def test_every_format_has_a_name_and_an_extension():
    assert (
        set(exporting.TITLES)
        == set(exporting.EXTENSIONS)
        == set(exporting.TABLE_FORMATS)
    )
    assert (
        exporting.EXTENSIONS["xlsx"] == ".xlsx" and exporting.EXTENSIONS["md"] == ".md"
    )


def test_an_export_says_what_it_is_and_how_big():
    assert PEOPLE.describe() == "Users of Sales (2 rows)"
    one = Exportable("Users of Sales", ("Name",), (("Ann",),))
    assert one.describe() == "Users of Sales (1 row)"
    assert (
        Exportable("Users of Sales", ("Name",)).describe() == "Users of Sales (0 rows)"
    )
    assert THING.describe() == "JSON of Sales" and THING.count == 1


def test_a_cell_that_shows_a_dash_is_empty_in_an_export():
    assert PEOPLE.cells[1] == ("Bob", "", "Viewer", "User")
    assert Exportable("t", ("a",), (("–",), (" - ",), ("-x",))).cells == [
        ("",),
        ("",),
        ("-x",),
    ]


def test_a_column_without_a_name_or_with_the_name_of_another_gets_one():
    odd = Exportable(
        "t", ("Fetched", "", "Name", "Name", ""), (("1", "2", "3", "4", "5"),)
    )

    assert odd.headers == ("Fetched", "column 2", "Name", "Name 2", "column 5")


# ---------------------------------------------------------------------------
# the text of each format
# ---------------------------------------------------------------------------


def test_csv_has_a_header_and_quotes_what_needs_it():
    item = Exportable(
        "t", ("Name", "Note"), (("Ann", 'says "hi", twice'), ("Bob", "two\nlines"))
    )

    text = exporting.to_text(item, "csv")

    assert text.startswith("Name,Note\r\n")
    assert list(csv.reader(io.StringIO(text, newline=""))) == [
        ["Name", "Note"],
        ["Ann", 'says "hi", twice'],
        ["Bob", "two\nlines"],
    ]


def test_a_cell_that_a_spreadsheet_would_run_is_written_as_text():
    item = Exportable(
        "t", ("A",), (("=HYPERLINK(1)",), ("+1",), ("-1",), ("@x",), ("ok",))
    )

    csv_rows = [r[0] for r in csv.reader(io.StringIO(exporting.to_text(item, "csv")))][
        1:
    ]
    tsv_rows = exporting.to_text(item, "tsv").splitlines()[1:]

    assert csv_rows == ["'=HYPERLINK(1)", "'+1", "'-1", "'@x", "ok"]
    assert tsv_rows == ["'=HYPERLINK(1)", "'+1", "'-1", "'@x", "ok"]
    assert "'" not in exporting.to_text(item, "json").replace('"ok"', "")


def test_a_header_that_a_spreadsheet_would_run_is_written_as_text_too():
    item = Exportable("t", ("=A", "+B", "ok"), (("x", "y", "z"),))

    assert exporting.to_text(item, "csv").splitlines()[0] == "'=A,'+B,ok"


def test_tab_separated_text_pastes_into_cells():
    item = Exportable("t", ("A", "B"), (("x\ty", "line\nbreak"), ("1", "2")))

    text = exporting.to_text(item, "tsv")

    assert text == "A\tB\nx y\tline break\n1\t2\n"
    assert exporting.clipboard_text(item) == text


def test_json_of_a_table_is_a_list_of_objects_that_keep_their_letters():
    found = json.loads(exporting.to_text(PEOPLE, "json"))

    assert found[0] == {
        "Name": "Ann Ås",
        "E-mail or id": "ann@example.com",
        "Access": "Admin",
        "Type": "User",
    }
    assert found[1]["E-mail or id"] == ""
    assert "Ås" in exporting.to_text(PEOPLE, "json")  # not Å


def test_json_of_other_data_is_the_data_and_is_what_the_clipboard_gets():
    text = exporting.to_text(THING, "json")

    assert json.loads(text) == THING.data and "Sales ü" in text and text.endswith("\n")
    assert exporting.clipboard_text(THING) == text


def test_a_markdown_table_has_its_header_rule_and_escapes_what_breaks_it():
    item = Exportable("t", ("Name", "Note"), (("A|B", "two\nlines"), ("C", "")))

    assert exporting.to_text(item, "md") == (
        "| Name | Note |\n| --- | --- |\n| A\\|B | two<br>lines |\n| C |  |\n"
    )


def test_the_clipboard_gets_a_table_as_tab_separated_text():
    assert exporting.clipboard_text(PEOPLE).splitlines()[:2] == [
        "Name\tE-mail or id\tAccess\tType",
        "Ann Ås\tann@example.com\tAdmin\tUser",
    ]


# ---------------------------------------------------------------------------
# files
# ---------------------------------------------------------------------------


def test_a_file_is_written_in_a_folder_that_is_made_when_it_is_not_there(tmp_path):
    target = tmp_path / "deep" / "er" / "people.json"

    written = exporting.write(PEOPLE, "json", target)

    assert (
        written == target and json.loads(target.read_text(encoding="utf-8"))[0]["Name"]
    )
    assert not target.read_bytes().startswith(b"\xef\xbb\xbf")  # no byte order mark


@pytest.mark.parametrize("fmt", ["csv", "tsv"])
def test_csv_and_tsv_files_start_with_a_byte_order_mark_so_that_excel_reads_the_letters(
    tmp_path, fmt
):
    target = exporting.write(PEOPLE, fmt, tmp_path / f"people.{fmt}")

    raw = target.read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf") and "Ann Ås".encode("utf-8") in raw
    if fmt == "csv":
        assert b"\r\n" in raw and b"\r\r" not in raw  # the line ends are not doubled


def test_a_workbook_has_a_bold_header_text_cells_and_a_name_that_excel_accepts(
    tmp_path,
):
    item = Exportable(
        "Users of [Confidential]IT: Management / Finance & more tail tail",
        ("Name", "Note"),
        (("=1+1", "wide " * 5), ("Ann", "")),
    )

    target = exporting.write(item, "xlsx", tmp_path / "people.xlsx")

    sheet = load_workbook(target).active
    assert len(sheet.title) <= 31 and not set("[]:*?/\\") & set(sheet.title)
    assert [c.value for c in sheet[1]] == ["Name", "Note"] and sheet["A1"].font.bold
    assert sheet["A2"].value == "=1+1" and sheet["A2"].data_type == "s"  # not a formula
    assert sheet["B3"].value in (None, "")
    assert sheet.freeze_panes == "A2" and sheet.auto_filter.ref == "A1:B3"
    assert sheet.column_dimensions["B"].width > sheet.column_dimensions["A"].width


def test_a_column_of_a_workbook_is_no_wider_than_sixty(tmp_path):
    item = Exportable("t", ("A", "B"), (("x" * 200, "y"),))

    sheet = load_workbook(exporting.write(item, "xlsx", tmp_path / "w.xlsx")).active

    assert sheet.column_dimensions["A"].width == 60
    assert sheet.column_dimensions["B"].width == 3  # the cell and two blanks


def test_an_empty_table_is_a_workbook_with_only_a_header(tmp_path):
    empty = Exportable("Nobody", ("Name",), ())

    target = exporting.write(empty, "xlsx", tmp_path / "nobody.xlsx")

    sheet = load_workbook(target).active
    assert sheet.max_row == 1 and sheet["A1"].value == "Name"
    assert exporting.to_text(empty, "csv") == "Name\r\n"


def test_the_name_of_a_worksheet_is_never_empty():
    assert exporting.sheet_name("[]:*?/\\") == "Export"
    assert exporting.sheet_name("x" * 100) == "x" * 31


# ---------------------------------------------------------------------------
# names
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    "title, stem",
    [
        ("Users of Sales", "Users of Sales"),
        (
            "Users of [Confidential]IT Management",
            "Users of [Confidential]IT Management",
        ),
        ('Users of a/b\\c:d*e?f"g<h>i|j', "Users of a-b-c-d-e-f-g-h-i-j"),
        ("  .Users.  ", "Users"),
        ("", "export"),
        ("///", "export"),
        ("CON", "_CON"),
        ("lpt1", "_lpt1"),
        ("nul.txt", "_nul.txt"),
        ("Users of Sales\nand\tmore", "Users of Sales and more"),
    ],
)
def test_a_title_becomes_a_name_that_every_system_accepts(title, stem):
    assert exporting.file_stem(title) == stem


def test_a_long_title_is_cut_without_leaving_a_dot_or_a_blank_at_the_end():
    stem = exporting.file_stem("A" * 78 + " .. more")

    assert len(stem) <= 80 and not stem.endswith((" ", ".", "-"))


def test_the_suggested_name_has_the_extension_of_the_format():
    assert exporting.file_name(PEOPLE, "csv") == "Users of Sales.csv"
    assert exporting.file_name(PEOPLE, "xlsx") == "Users of Sales.xlsx"
    assert exporting.file_name(THING, "json") == "JSON of Sales.json"


def test_a_name_that_is_taken_gets_a_number_before_its_extension(tmp_path):
    assert exporting.free_path(tmp_path, "a.csv") == tmp_path / "a.csv"
    (tmp_path / "a.csv").write_text("x")
    assert exporting.free_path(tmp_path, "a.csv") == tmp_path / "a (2).csv"
    (tmp_path / "a (2).csv").write_text("x")
    assert exporting.free_path(tmp_path, "a.csv") == tmp_path / "a (3).csv"


def test_rows_of_anything_that_shows_as_text_are_tuples_of_text():
    assert exporting.rows_of([[1, None, "x"], (2.5, "y", True)]) == (
        ("1", "", "x"),
        ("2.5", "y", "True"),
    )
